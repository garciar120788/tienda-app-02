import csv
import io
import json
import os
import re
import tempfile
import time
import uuid

import psycopg
from flask import (Blueprint, Response, abort, flash, redirect, render_template, request,
                   send_file, session, url_for)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .db import get_db
from .utils import ErrorValidacion, a_codigo, a_decimal, a_entero, sin_acentos

bp = Blueprint("inventario", __name__, url_prefix="/inventario")

POR_PAGINA = 50
DIR_IMPORT = os.path.join(tempfile.gettempdir(), f"tienda_import_{os.getuid()}")

COLUMNAS = ["codigo", "nombre", "precio", "costo", "stock", "stock_minimo"]
ALIAS = {
    "codigo": "codigo", "sku": "codigo", "codigo de barras": "codigo", "clave": "codigo",
    "nombre": "nombre", "producto": "nombre", "descripcion": "nombre",
    "precio": "precio", "precio venta": "precio", "precio de venta": "precio",
    "costo": "costo", "precio compra": "costo", "precio de compra": "costo",
    "stock": "stock", "existencia": "stock", "existencias": "stock", "cantidad": "stock",
    "stock minimo": "stock_minimo", "stock_minimo": "stock_minimo", "minimo": "stock_minimo",
}


# ======================= Utilidades =======================

def _producto_o_404(conn, pid, bloquear=False):
    sql = "SELECT * FROM productos WHERE id = %s" + (" FOR UPDATE" if bloquear else "")
    p = conn.execute(sql, (pid,)).fetchone()
    if not p:
        abort(404)
    return p


def _leer_form_producto(form, nuevo):
    datos = {
        "codigo": a_codigo(form.get("codigo")),
        "nombre": (form.get("nombre") or "").strip(),
        "precio": a_decimal(form.get("precio"), "Precio"),
        "costo": a_decimal(form.get("costo"), "Costo", requerido=False) or 0,
        "stock_minimo": a_entero(form.get("stock_minimo"), "Stock mínimo", requerido=False) or 0,
        "es_servicio": form.get("es_servicio") == "1",
    }
    if datos["es_servicio"]:
        datos["stock_minimo"] = 0
    if nuevo:
        datos["stock"] = 0 if datos["es_servicio"] else (
            a_entero(form.get("stock"), "Stock inicial", requerido=False) or 0)
    if not datos["codigo"]:
        raise ErrorValidacion("El código es obligatorio")
    if len(datos["codigo"]) > 64:
        raise ErrorValidacion("El código es demasiado largo (máx. 64 caracteres)")
    if not datos["nombre"]:
        raise ErrorValidacion("El nombre es obligatorio")
    if len(datos["nombre"]) > 200:
        raise ErrorValidacion("El nombre es demasiado largo (máx. 200 caracteres)")
    return datos


def registrar_movimiento(conn, producto_id, tipo, cantidad, stock_resultante, venta_id=None, nota=None):
    conn.execute(
        """INSERT INTO movimientos (producto_id, tipo, cantidad, stock_resultante, venta_id, nota)
           VALUES (%s, %s, %s, %s, %s, %s)""",
        (producto_id, tipo, cantidad, stock_resultante, venta_id, nota),
    )


# ======================= Listado y CRUD =======================

@bp.get("/")
def lista():
    conn = get_db()
    q = request.args.get("q", "").strip()
    filtro = request.args.get("filtro", "activos")
    pagina = max(request.args.get("pagina", 1, type=int), 1)

    condiciones, params = [], []
    if filtro == "bajo":
        condiciones.append("activo AND NOT es_servicio AND stock <= stock_minimo")
    elif filtro == "inactivos":
        condiciones.append("NOT activo")
    elif filtro != "todos":
        filtro = "activos"
        condiciones.append("activo")
    if q:
        condiciones.append("(codigo ILIKE %s OR nombre ILIKE %s)")
        like = f"%{q}%"
        params += [like, like]
    where = ("WHERE " + " AND ".join(condiciones)) if condiciones else ""

    total = conn.execute(f"SELECT count(*) AS n FROM productos {where}", params).fetchone()["n"]
    productos = conn.execute(
        f"SELECT * FROM productos {where} ORDER BY nombre LIMIT %s OFFSET %s",
        params + [POR_PAGINA, (pagina - 1) * POR_PAGINA],
    ).fetchall()
    paginas = max((total + POR_PAGINA - 1) // POR_PAGINA, 1)
    return render_template(
        "inventario/lista.html", productos=productos, q=q, filtro=filtro,
        pagina=pagina, paginas=paginas, total=total,
    )


@bp.route("/nuevo", methods=["GET", "POST"])
def nuevo():
    form = {}
    if request.method == "POST":
        form = request.form
        conn = get_db()
        try:
            d = _leer_form_producto(request.form, nuevo=True)
            p = conn.execute(
                """INSERT INTO productos (codigo, nombre, precio, costo, stock, stock_minimo, es_servicio)
                   VALUES (%(codigo)s, %(nombre)s, %(precio)s, %(costo)s, %(stock)s, %(stock_minimo)s,
                           %(es_servicio)s)
                   RETURNING id""",
                d,
            ).fetchone()
            if not d["es_servicio"]:
                registrar_movimiento(conn, p["id"], "alta", d["stock"], d["stock"], nota="Alta de producto")
            conn.commit()
            flash(f"Producto «{d['nombre']}» creado.", "success")
            if request.form.get("otro"):
                return redirect(url_for("inventario.nuevo"))
            return redirect(url_for("inventario.detalle", pid=p["id"]))
        except ErrorValidacion as e:
            conn.rollback()
            flash(str(e), "danger")
        except psycopg.errors.UniqueViolation:
            conn.rollback()
            flash("Ya existe un producto con ese código.", "danger")
    return render_template("inventario/form.html", producto=None, form=form)


@bp.get("/<int:pid>")
def detalle(pid):
    conn = get_db()
    p = _producto_o_404(conn, pid)
    movimientos = conn.execute(
        "SELECT * FROM movimientos WHERE producto_id = %s ORDER BY fecha DESC, id DESC LIMIT 200", (pid,)
    ).fetchall()
    vendidos = conn.execute(
        """SELECT COALESCE(sum(d.cantidad), 0) AS piezas, COALESCE(sum(d.subtotal), 0) AS total
           FROM venta_detalle d JOIN ventas v ON v.id = d.venta_id
           WHERE d.producto_id = %s AND v.estado = 'completada'
             AND v.fecha >= current_date - 29""",
        (pid,),
    ).fetchone()
    return render_template("inventario/detalle.html", p=p, movimientos=movimientos, vendidos=vendidos)


@bp.route("/<int:pid>/editar", methods=["GET", "POST"])
def editar(pid):
    conn = get_db()
    p = _producto_o_404(conn, pid)
    form = p
    if request.method == "POST":
        form = request.form
        try:
            d = _leer_form_producto(request.form, nuevo=False)
            d["id"] = pid
            conn.execute(
                """UPDATE productos SET codigo = %(codigo)s, nombre = %(nombre)s, precio = %(precio)s,
                          costo = %(costo)s, stock_minimo = %(stock_minimo)s,
                          es_servicio = %(es_servicio)s, actualizado = now()
                   WHERE id = %(id)s""",
                d,
            )
            conn.commit()
            flash("Producto actualizado.", "success")
            return redirect(url_for("inventario.detalle", pid=pid))
        except ErrorValidacion as e:
            conn.rollback()
            flash(str(e), "danger")
        except psycopg.errors.UniqueViolation:
            conn.rollback()
            flash("Ya existe otro producto con ese código.", "danger")
    return render_template("inventario/form.html", producto=p, form=form)


@bp.post("/<int:pid>/stock")
def ajustar_stock(pid):
    conn = get_db()
    tipo = request.form.get("tipo")
    nota = (request.form.get("nota") or "").strip()[:300] or None
    try:
        p = _producto_o_404(conn, pid, bloquear=True)
        if p["es_servicio"]:
            raise ErrorValidacion("Este producto es un servicio y no maneja stock")
        if tipo == "entrada":
            cantidad = a_entero(request.form.get("cantidad"), "Cantidad", minimo=1)
            nuevo_stock = p["stock"] + cantidad
        elif tipo == "ajuste":
            nuevo_stock = a_entero(request.form.get("cantidad"), "Stock real", minimo=0)
            cantidad = nuevo_stock - p["stock"]
            if cantidad == 0:
                conn.rollback()
                flash("El stock ya tenía ese valor; no se registró ningún cambio.", "info")
                return redirect(url_for("inventario.detalle", pid=pid))
            if not nota:
                raise ErrorValidacion("Escribe un motivo para el ajuste (ej. conteo físico, merma, daño)")
        else:
            raise ErrorValidacion("Tipo de movimiento inválido")
        conn.execute(
            "UPDATE productos SET stock = %s, actualizado = now() WHERE id = %s", (nuevo_stock, pid)
        )
        registrar_movimiento(conn, pid, tipo, cantidad, nuevo_stock, nota=nota)
        conn.commit()
        flash(f"Stock actualizado: {p['stock']} → {nuevo_stock}.", "success")
    except ErrorValidacion as e:
        conn.rollback()
        flash(str(e), "danger")
    return redirect(url_for("inventario.detalle", pid=pid))


@bp.post("/<int:pid>/activo")
def cambiar_activo(pid):
    conn = get_db()
    p = _producto_o_404(conn, pid)
    conn.execute("UPDATE productos SET activo = NOT activo, actualizado = now() WHERE id = %s", (pid,))
    conn.commit()
    flash("Producto reactivado." if not p["activo"] else
          "Producto desactivado: ya no aparecerá en la caja.", "success")
    return redirect(url_for("inventario.detalle", pid=pid))


@bp.get("/exportar.csv")
def exportar():
    conn = get_db()
    filas = conn.execute(
        "SELECT codigo, nombre, precio, costo, stock, stock_minimo, activo FROM productos ORDER BY nombre"
    ).fetchall()
    out = io.StringIO()
    out.write("﻿")  # BOM para que Excel respete acentos
    w = csv.writer(out)
    w.writerow(COLUMNAS + ["activo"])
    for f in filas:
        w.writerow([f["codigo"], f["nombre"], f["precio"], f["costo"], f["stock"], f["stock_minimo"],
                    "si" if f["activo"] else "no"])
    return Response(out.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename=inventario.csv"})


# ======================= Importación Excel / CSV =======================

@bp.get("/plantilla.<fmt>")
def plantilla(fmt):
    ejemplo = [
        ["7501234567890", "Refresco cola 600 ml", 18.50, 12.00, 24, 6],
        ["SKU-0001", "Galletas surtidas", 25.00, 17.50, 10, 3],
    ]
    if fmt == "csv":
        out = io.StringIO()
        out.write("﻿")
        w = csv.writer(out)
        w.writerow(COLUMNAS)
        w.writerows(ejemplo)
        return Response(out.getvalue(), mimetype="text/csv; charset=utf-8",
                        headers={"Content-Disposition": "attachment; filename=plantilla_inventario.csv"})
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Inventario"
        ws.append(COLUMNAS)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0D6EFD")
        for fila in ejemplo:
            ws.append(fila)
        # Columna de código como texto para no perder ceros a la izquierda
        for fila in ws.iter_rows(min_row=1, max_row=1000, min_col=1, max_col=1):
            for c in fila:
                c.number_format = "@"
        for col, ancho in zip("ABCDEF", [18, 34, 10, 10, 10, 14]):
            ws.column_dimensions[col].width = ancho
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name="plantilla_inventario.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    abort(404)


def _normalizar_encabezado(h):
    h = sin_acentos(str(h or "")).strip().lower().replace("_", " ")
    h = re.sub(r"\s+", " ", h)
    return ALIAS.get(h) or ALIAS.get(h.replace(" ", "_"))


def _leer_archivo(archivo):
    """Devuelve una lista de filas (listas) incluyendo el encabezado."""
    nombre = (archivo.filename or "").lower()
    contenido = archivo.read()
    if nombre.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception:
            raise ErrorValidacion("No se pudo leer el archivo de Excel. Verifica que sea .xlsx válido.")
        ws = wb.worksheets[0]
        return [list(f) for f in ws.iter_rows(values_only=True)]
    if nombre.endswith(".csv") or nombre.endswith(".txt"):
        for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                texto = contenido.decode(codificacion)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
        except csv.Error:
            dialecto = csv.excel
        return list(csv.reader(io.StringIO(texto), dialecto))
    raise ErrorValidacion("Formato no soportado. Sube un archivo .xlsx o .csv")


def analizar_filas(conn, filas):
    """Valida las filas del archivo y determina qué se creará o actualizará."""
    # Saltar filas vacías al inicio
    while filas and all(c is None or str(c).strip() == "" for c in filas[0]):
        filas.pop(0)
    if not filas:
        raise ErrorValidacion("El archivo está vacío")

    encabezado = [_normalizar_encabezado(h) for h in filas[0]]
    indices = {}
    for i, campo in enumerate(encabezado):
        if campo and campo not in indices:
            indices[campo] = i
    faltantes = [c for c in ("codigo", "nombre", "precio") if c not in indices]
    if faltantes:
        raise ErrorValidacion(
            "Faltan columnas obligatorias: " + ", ".join(faltantes)
            + ". Usa la plantilla como guía (codigo, nombre, precio, costo, stock, stock_minimo)."
        )

    existentes = {
        r["codigo"]: r for r in conn.execute("SELECT id, codigo, nombre, stock FROM productos").fetchall()
    }
    resultado, vistos = [], set()
    for n, fila in enumerate(filas[1:], start=2):
        if all(c is None or str(c).strip() == "" for c in fila):
            continue

        def celda(campo):
            i = indices.get(campo)
            return fila[i] if i is not None and i < len(fila) else None

        item = {"fila": n, "errores": []}
        try:
            item["codigo"] = a_codigo(celda("codigo"))
            if not item["codigo"]:
                raise ErrorValidacion("Código vacío")
            if len(item["codigo"]) > 64:
                raise ErrorValidacion("Código demasiado largo")
            if item["codigo"] in vistos:
                raise ErrorValidacion("Código repetido en el archivo")
            vistos.add(item["codigo"])
            item["nombre"] = str(celda("nombre") or "").strip()[:200]
            if not item["nombre"]:
                raise ErrorValidacion("Nombre vacío")
            item["precio"] = str(a_decimal(celda("precio"), "Precio"))
            costo = a_decimal(celda("costo"), "Costo", requerido=False)
            item["costo"] = str(costo) if costo is not None else None
            item["stock"] = a_entero(celda("stock"), "Stock", requerido=False)
            item["stock_minimo"] = a_entero(celda("stock_minimo"), "Stock mínimo", requerido=False)
        except ErrorValidacion as e:
            item["errores"].append(str(e))

        existente = existentes.get(item.get("codigo"))
        item["accion"] = "actualizar" if existente else "nuevo"
        item["stock_actual"] = existente["stock"] if existente else None
        resultado.append(item)

    if not resultado:
        raise ErrorValidacion("El archivo no contiene productos")
    if len(resultado) > 10000:
        raise ErrorValidacion("El archivo tiene más de 10,000 filas; divídelo en partes")
    return resultado


def _limpiar_imports_viejos(horas=6):
    limite = time.time() - horas * 3600
    for nombre in os.listdir(DIR_IMPORT):
        ruta = os.path.join(DIR_IMPORT, nombre)
        try:
            if os.path.getmtime(ruta) < limite:
                os.remove(ruta)
        except OSError:
            pass


def _ruta_import(token):
    if not token or not re.fullmatch(r"[0-9a-f]{32}", token):
        return None
    return os.path.join(DIR_IMPORT, token + ".json")


@bp.route("/importar", methods=["GET", "POST"])
def importar():
    if request.method == "GET":
        return render_template("inventario/importar.html")

    archivo = request.files.get("archivo")
    if not archivo or not archivo.filename:
        flash("Selecciona un archivo.", "danger")
        return redirect(url_for("inventario.importar"))
    try:
        filas = analizar_filas(get_db(), _leer_archivo(archivo))
    except ErrorValidacion as e:
        flash(str(e), "danger")
        return redirect(url_for("inventario.importar"))

    os.makedirs(DIR_IMPORT, exist_ok=True)
    _limpiar_imports_viejos()
    token = uuid.uuid4().hex
    with open(_ruta_import(token), "w", encoding="utf-8") as f:
        json.dump(filas, f)
    session["import_token"] = token

    resumen = {
        "nuevos": sum(1 for f in filas if not f["errores"] and f["accion"] == "nuevo"),
        "actualizar": sum(1 for f in filas if not f["errores"] and f["accion"] == "actualizar"),
        "errores": sum(1 for f in filas if f["errores"]),
    }
    return render_template("inventario/importar_vista.html", filas=filas, resumen=resumen,
                           nombre_archivo=archivo.filename)


@bp.post("/importar/confirmar")
def importar_confirmar():
    token = session.pop("import_token", None)
    ruta = _ruta_import(token)
    if not ruta or not os.path.exists(ruta):
        flash("La vista previa expiró. Sube el archivo de nuevo.", "warning")
        return redirect(url_for("inventario.importar"))
    with open(ruta, encoding="utf-8") as f:
        filas = json.load(f)
    os.remove(ruta)

    modo = request.form.get("modo_stock", "sumar")
    conn = get_db()
    creados = actualizados = 0
    try:
        for it in filas:
            if it["errores"]:
                continue
            actual = conn.execute(
                "SELECT * FROM productos WHERE codigo = %s FOR UPDATE", (it["codigo"],)
            ).fetchone()
            if actual is None:
                stock = it["stock"] or 0
                nuevo = conn.execute(
                    """INSERT INTO productos (codigo, nombre, precio, costo, stock, stock_minimo)
                       VALUES (%s, %s, %s, %s, %s, %s) RETURNING id""",
                    (it["codigo"], it["nombre"], it["precio"], it["costo"] or 0, stock,
                     it["stock_minimo"] or 0),
                ).fetchone()
                registrar_movimiento(conn, nuevo["id"], "importacion", stock, stock,
                                     nota="Alta por importación")
                creados += 1
            else:
                if it["stock"] is None or actual["es_servicio"]:
                    stock = actual["stock"]
                elif modo == "reemplazar":
                    stock = it["stock"]
                else:
                    stock = actual["stock"] + it["stock"]
                conn.execute(
                    """UPDATE productos SET nombre = %s, precio = %s,
                              costo = COALESCE(%s, costo),
                              stock_minimo = COALESCE(%s, stock_minimo),
                              stock = %s, activo = TRUE, actualizado = now()
                       WHERE id = %s""",
                    (it["nombre"], it["precio"], it["costo"], it["stock_minimo"], stock, actual["id"]),
                )
                delta = stock - actual["stock"]
                if delta != 0:
                    nota = "Importación (reemplazo de stock)" if modo == "reemplazar" else "Importación (entrada)"
                    registrar_movimiento(conn, actual["id"], "importacion", delta, stock, nota=nota)
                actualizados += 1
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    flash(f"Importación completada: {creados} productos nuevos, {actualizados} actualizados.", "success")
    return redirect(url_for("inventario.lista"))
